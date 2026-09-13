"""Create the main-experiment workbook from append-only seed records."""

from __future__ import annotations

import math
from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Any, Iterable

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DATASETS, MODEL_NAME, RESULT_COLUMNS
from .records import latest_records


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
BODY_FONT = Font(name="Calibri", size=11, color="000000")


def _number(row: dict[str, Any], key: str) -> float | None:
    value = row.get(key)
    if value in (None, "", "N/A"):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _mean_std(values: list[float], decimals: int) -> str:
    if not values:
        return "N/A"
    mean = float(np.mean(values))
    std = float(np.std(values, ddof=0))
    return f"{mean:.{decimals}f}+/-{std:.{decimals}f}"


def _aggregate_group(rows: list[dict[str, Any]], planned_seed_count: int) -> tuple[list[Any], dict[str, float]]:
    ok_rows = [row for row in rows if str(row.get("status", "")).upper() == "OK"]
    statuses = {str(row.get("status", "")).upper() for row in rows}
    seeds_cell = f"{len(ok_rows)}/{planned_seed_count}"
    if not rows:
        return [seeds_cell, "N/A", "N/A", "N/A", "N/A", "N/A", "NOT RUN"], {}
    if "OOM" in statuses or "SKIPPED_OOM" in statuses:
        return [seeds_cell, "OOM", "OOM", "OOM", "OOM", "N/A", "OOM"], {}
    if not ok_rows:
        status = "FAILED" if "FAILED" in statuses else sorted(statuses)[0]
        return [seeds_cell, "N/A", "N/A", "N/A", "N/A", "N/A", status], {}

    metric_values = {
        "hits1": [value for row in ok_rows if (value := _number(row, "hits1")) is not None],
        "hits10": [value for row in ok_rows if (value := _number(row, "hits10")) is not None],
        "mrr": [value for row in ok_rows if (value := _number(row, "mrr")) is not None],
        "time_s": [value for row in ok_rows if (value := _number(row, "time_s")) is not None],
        "memory_gb": [value for row in ok_rows if (value := _number(row, "memory_gb")) is not None],
    }
    status = "OK" if len(ok_rows) == planned_seed_count else "PARTIAL"
    cells = [
        seeds_cell,
        _mean_std(metric_values["hits1"], 4),
        _mean_std(metric_values["hits10"], 4),
        _mean_std(metric_values["mrr"], 4),
        _mean_std(metric_values["time_s"], 2),
        _mean_std(metric_values["memory_gb"], 4),
        status,
    ]
    means = {
        key: float(np.mean(values))
        for key, values in metric_values.items()
        if values
    }
    return cells, means


def write_results_workbook(
    records: Iterable[dict[str, Any]],
    output_path: str | Path,
    *,
    planned_seeds: Iterable[int],
) -> Path:
    """Write one sheet per retained dataset in the canonical result format."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    planned_seed_list = [int(seed) for seed in planned_seeds]
    if not planned_seed_list:
        raise ValueError("planned_seeds cannot be empty")

    latest = latest_records(records)
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for (dataset, model, seed), row in latest.items():
        if seed in planned_seed_list:
            grouped[(dataset, model)].append(row)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for dataset in DATASETS:
        sheet = workbook.create_sheet(dataset)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:H2"
        for column, header in enumerate(RESULT_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.fill = copy(HEADER_FILL)
            cell.font = copy(HEADER_FONT)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 22

        cells, _ = _aggregate_group(
            grouped.get((dataset, MODEL_NAME), []),
            len(planned_seed_list),
        )
        for column, value in enumerate([MODEL_NAME, *cells], start=1):
            cell = sheet.cell(row=2, column=column, value=value)
            cell.font = copy(BODY_FONT)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center",
                vertical="center",
            )

        widths = (20, 10, 19, 19, 19, 16, 18, 14)
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width

    workbook.save(output_path)
    workbook.close()
    return output_path
